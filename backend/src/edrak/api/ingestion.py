"""Document parsing, RAG indexing, and profile ingestion for EDRAK."""

from __future__ import annotations

import io
import json
import logging
from typing import Any, Optional

import aiosqlite

from edrak.core.config import settings
from edrak.db.repository import add_company_document, save_company_profile

logger = logging.getLogger(__name__)


def extract_document_text(filename: str, content: bytes) -> tuple[str, dict[str, Any]]:
    """Extract plain text and size metrics from various file formats."""
    lower_name = filename.lower()
    size_kb = round(len(content) / 1024, 1)

    # PDF format
    if lower_name.endswith(".pdf"):
        try:
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(content))
            pages_text: list[str] = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages_text.append(text)
            full_text = "\n\n".join(pages_text)
            page_count = len(reader.pages)
            return full_text, {"unit": "pages", "value": page_count}
        except Exception as exc:
            logger.warning("PDF extraction failed for %s: %s", filename, exc)
            return "", {"unit": "kb", "value": size_kb}

    # Word DOCX format
    if lower_name.endswith(".docx"):
        try:
            import docx

            doc = docx.Document(io.BytesIO(content))
            full_text = "\n".join(p.text for p in doc.paragraphs if p.text)
            return full_text, {"unit": "pages", "value": max(1, len(doc.paragraphs) // 10)}
        except Exception as exc:
            logger.warning("DOCX extraction failed for %s: %s", filename, exc)
            return "", {"unit": "kb", "value": size_kb}

    # Markdown and plain text
    if lower_name.endswith((".md", ".txt")):
        text = content.decode("utf-8", errors="replace")
        line_count = len(text.splitlines())
        return text, {"unit": "rows", "value": line_count}

    # JSON format
    if lower_name.endswith(".json"):
        text = content.decode("utf-8", errors="replace")
        return text, {"unit": "kb", "value": size_kb}

    # Fallback
    text = content.decode("utf-8", errors="replace")
    return text, {"unit": "kb", "value": size_kb}


def is_company_profile_json(data: Any) -> bool:
    """Check if parsed JSON data matches an EDRAK company profile structure."""
    if not isinstance(data, dict):
        return False
    if "name" not in data:
        return False
    profile_keys = (
        "products",
        "offerings",
        "aliases",
        "description",
        "notes",
        "legal_name",
        "business_model",
        "strategic_priorities",
        "differentiators",
    )
    return any(k in data for k in profile_keys)


async def ingest_company_profile_data(
    data: dict[str, Any], workspace_id: str, db: aiosqlite.Connection
) -> None:
    """Save full profile to disk and map fields to SQLite for UI synchronization."""
    active_profile_path = settings.BASE_DIR / "data" / "profiles" / "active_profile.json"
    active_profile_path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Save full JSON on disk (preserving deep/rich nested fields)
    active_profile_path.write_text(
        json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    # 2. Extract and adapt fields for the UI form
    name = str(data.get("name") or "Company").strip()
    aliases = list(data.get("aliases") or [])
    industry = data.get("industry")
    description = str(data.get("description") or data.get("notes") or "").strip()

    offerings = list(data.get("offerings") or data.get("products") or [])
    markets = list(data.get("markets") or [])
    if not markets and isinstance(data.get("tam"), dict):
        core = data["tam"].get("core_market")
        if core:
            markets.append(str(core))

    strategic_goals = str(data.get("strategic_goals") or "")
    if not strategic_goals and isinstance(data.get("strategic_priorities"), list):
        strategic_goals = "; ".join(str(p) for p in data["strategic_priorities"])

    website = str(data.get("website") or "")
    socials = list(data.get("socials") or [])

    fields_dict = {
        "name": name,
        "aliases": aliases,
        "industry": industry,
        "description": description,
        "offerings": offerings,
        "markets": markets,
        "strategic_goals": strategic_goals,
        "website": website,
        "socials": socials,
    }

    await save_company_profile(db, workspace_id, fields_dict, completed=True)
    logger.info("Successfully ingested active company profile: %s", name)


def index_text_into_rag(filename: str, text: str, workspace_id: str) -> bool:
    """Chunk and index document text into ChromaDB vector store."""
    if not text.strip():
        return False
    try:
        from edrak.rag.indexer import InternalIndexer

        indexer = InternalIndexer()
        count = indexer.index_document(
            content=text,
            source_uri=f"internal://{filename}",
            title=filename,
            doc_type="internal_doc",
            extra_metadata={"filename": filename},
            is_synthetic=False,
            fallback_title=filename,
            workspace_id=workspace_id,
        )
        logger.info("Indexed %d chunks for document: %s", count, filename)
        return count > 0
    except Exception as exc:
        logger.warning("Could not index document %s into RAG: %s", filename, exc)
        return False


def _sync_index_state(workspace_id: str) -> list[dict[str, Any]]:
    """Blocking part of the startup sync: tag old chunks and list what the demo workspace has indexed."""
    from edrak.rag.indexer import InternalIndexer

    indexer = InternalIndexer()
    tagged = indexer.backfill_workspace(workspace_id)
    if tagged:
        logger.info("Assigned %d untagged chunks to workspace %s", tagged, workspace_id)
    return indexer.workspace_documents(workspace_id)


async def sync_demo_workspace(db: aiosqlite.Connection) -> None:
    """List the pre-loaded knowledge (data/internal, the handbook) as indexed documents.

    The documents come from the vector index itself, so a file shows as indexed only if its
    chunks are there. Rows get fixed ids, which makes this safe to run on every start.
    """
    import asyncio
    import re

    workspace_id = settings.DEMO_WORKSPACE_ID
    cursor = await db.execute("SELECT 1 FROM workspaces WHERE workspace_id = ?;", (workspace_id,))
    if not await cursor.fetchone():
        return
    try:
        indexed = await asyncio.to_thread(_sync_index_state, workspace_id)
    except Exception as exc:  # noqa: BLE001 - the app must still start without the index
        logger.warning("Could not read the knowledge index: %s", exc)
        return

    cursor = await db.execute(
        "SELECT filename FROM company_documents WHERE workspace_id = ?;", (workspace_id,)
    )
    listed = {row["filename"] for row in await cursor.fetchall()}
    for doc in indexed:
        if doc["kind"] == "handbook":
            doc_id = "doc-demo-handbook"
            name = f"GitLab handbook ({doc['chunks']} pages)"
            size = {"unit": "pages", "value": doc["chunks"]}
            stale = "GitLab handbook (%"
        else:
            doc_id = "doc-demo-" + re.sub(r"[^a-z0-9]+", "-", doc["filename"].lower()).strip("-")
            name = doc["filename"]
            on_disk = next(settings.get_absolute_internal_data_path().rglob(doc["filename"]), None)
            size = {"unit": "kb", "value": round(on_disk.stat().st_size / 1024, 1)} if on_disk else None
            stale = None
        if stale:
            await db.execute(
                "DELETE FROM company_documents WHERE document_id = ? OR (workspace_id = ? AND filename LIKE ?);",
                (doc_id, workspace_id, stale),
            )
        elif name in listed:
            continue  # already there, for example uploaded through the UI
        await add_company_document(
            db, doc_id, workspace_id, name, size, status="indexed"
        )
